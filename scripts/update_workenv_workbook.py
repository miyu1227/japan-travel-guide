from copy import copy
from datetime import datetime

import openpyxl


PATH = r"C:\Users\miyul\Downloads\作業環境測定士.xlsx"
BLOCK_SHEETS = [
    "① 労働衛生一般（共通）",
    "② 労働衛生関係法令（共通）",
    "③ デザイン・サンプリング（共通）",
    "④ 分析に関する概論（共通）",
    "⑤ 有機溶剤（第1種）",
    "⑥ 鉱物性粉じん（第1種）",
    "⑦ 特定化学物質（第1種）",
    "⑧ 金属類（第1種）",
    "⑨ 放射性物質（第1種）",
]
HEADERS = ["id", "block", "question", "a", "b", "c", "d", "answer", "explanation"]


def set_cell(ws, row, col_name, value):
    col = HEADERS.index(col_name) + 1
    ws.cell(row=row, column=col).value = value


wb = openpyxl.load_workbook(PATH)

repairs = []

ws = wb["① 労働衛生一般（共通）"]
set_cell(
    ws,
    39,
    "explanation",
    "日射がない屋内環境のWBGTは、0.7×自然湿球温度＋0.3×黒球温度で算出します。",
)
repairs.append("①-38 WBGT解説に係数を明示")

ws = wb["② 労働衛生関係法令（共通）"]
set_cell(
    ws,
    34,
    "a",
    "当該作業が行われる時間のうち、測定対象物質の濃度が最も高くなると思われる連続した15分間",
)
set_cell(
    ws,
    34,
    "explanation",
    "D測定（B測定に相当）は、濃度が最も高くなると思われる時間に行い、その試料空気の採取等の時間は連続した15分間です。",
)
repairs.append("②-33 D測定の採取時間を連続15分間として明確化")

final_ws = wb["最終フォーマット"]
for row in final_ws.iter_rows():
    for cell in row:
        cell.value = None

for c, header in enumerate(HEADERS, start=1):
    final_ws.cell(row=1, column=c).value = header

if BLOCK_SHEETS:
    template_ws = wb[BLOCK_SHEETS[0]]
    for c in range(1, len(HEADERS) + 1):
        src = template_ws.cell(row=1, column=c)
        dst = final_ws.cell(row=1, column=c)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)

out_row = 2
global_id = 1
for sheet_name in BLOCK_SHEETS:
    ws = wb[sheet_name]
    for row in range(2, ws.max_row + 1):
        if not ws.cell(row=row, column=3).value:
            continue
        values = [ws.cell(row=row, column=c).value for c in range(1, len(HEADERS) + 1)]
        values[0] = global_id
        for c, value in enumerate(values, start=1):
            cell = final_ws.cell(row=out_row, column=c)
            cell.value = value
            src = ws.cell(row=row, column=c)
            if src.has_style:
                cell._style = copy(src._style)
            if src.number_format:
                cell.number_format = src.number_format
            if src.alignment:
                cell.alignment = copy(src.alignment)
        out_row += 1
        global_id += 1

for c in range(1, len(HEADERS) + 1):
    final_ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = max(
        final_ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width or 10,
        12 if c in [1, 2, 8] else 36,
    )

wb.save(PATH)
print({"saved": PATH, "repairs": repairs, "final_rows": out_row - 1, "timestamp": datetime.now().isoformat()})
