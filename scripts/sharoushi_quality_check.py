import json
import re
import unicodedata
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl


XLSX_PATH = Path(r"C:\Users\miyul\Downloads\社会保険労務士.xlsx")
REPORT_PATH = Path(r"C:\tmp\sharoushi_quality_report.json")
SHEET_PREFIXES = tuple("①②③④⑤⑥⑦⑧⑨⑩")
EXPECTED_HEADERS = ["id", "block", "question", "a", "b", "c", "d", "answer", "explanation"]

AMBIGUOUS_TERMS = ["最も適切", "一般的に", "主に", "通常"]
NG_REPLACEMENTS = [
    ("絶対的明示事項", "明示が必要な事項"),
    ("絶対的支給制限事由", "保険給付を行わない事由"),
    ("合格保証", "合格を約束するもの"),
    ("100%", "すべて"),
    ("絶対", "原則として"),
    ("必ず", "原則として"),
    ("確実", "適切"),
    ("保証", "担保"),
]
REF_REPLACEMENTS = [
    ("テキストでは", "この論点では"),
    ("本書では", "この問題では"),
    ("本書に", "この問題に"),
    ("下表", "該当する一覧"),
    ("上図", "該当する図"),
]
APP_LIMITS = {
    "アプリ名": 30,
    "サブタイトル": 30,
    "プロモーション用テキスト": 170,
    "概要": 4000,
    "キーワード": 100,
}
APP_REQUIRED = [
    "アプリ名",
    "サブタイトル",
    "プロモーション用テキスト",
    "概要",
    "キーワード",
    "スクリーンショット① メインビジュアル",
    "スクリーンショット② 網羅性訴求",
    "スクリーンショット③ 復習機能",
    "スクリーンショット④ 模試モード",
    "スクリーンショット⑤ 習熟度マップ",
]


def cell_text(value):
    return "" if value is None else str(value)


def normalize_answer(value):
    text = unicodedata.normalize("NFKC", cell_text(value)).strip().lower()
    return text


def norm_question(value):
    return cell_text(value).strip()


def question_prefix(value, n=30):
    return norm_question(value)[:n]


def row_has_question(ws, row):
    return norm_question(ws.cell(row, 3).value) != ""


def data_rows(ws):
    return [r for r in range(2, ws.max_row + 1) if any(ws.cell(r, c).value is not None for c in range(1, 10))]


def compact_ids(ws):
    next_id = 1
    for r in range(2, ws.max_row + 1):
        if any(ws.cell(r, c).value is not None for c in range(1, 10)):
            ws.cell(r, 1).value = next_id
            next_id += 1


def append_issue(report, key, sheet, row, item_id, detail):
    report[key].append({"sheet": sheet, "row": row, "id": item_id, "detail": detail})


def inspect_block_sheet(wb, sheet_name, report):
    ws = wb[sheet_name]
    before = len(data_rows(ws))
    report["block_summary"][sheet_name] = {
        "before": before,
        "after": before,
        "auto_fixes": 0,
        "needs_review": 0,
    }

    duplicate_rows = []
    seen_questions = {}

    for r in range(2, ws.max_row + 1):
        if not any(ws.cell(r, c).value is not None for c in range(1, 10)):
            continue

        item_id = ws.cell(r, 1).value
        question = norm_question(ws.cell(r, 3).value)
        choices = [cell_text(ws.cell(r, c).value) for c in range(4, 8)]
        answer = normalize_answer(ws.cell(r, 8).value)
        explanation = cell_text(ws.cell(r, 9).value)

        if answer != cell_text(ws.cell(r, 8).value):
            old = ws.cell(r, 8).value
            ws.cell(r, 8).value = answer
            report["auto_fixes"].append({"sheet": sheet_name, "id": item_id, "kind": "answer正規化", "before": old, "after": answer})

        if answer not in {"a", "b", "c", "d"}:
            append_issue(report, "needs_review", sheet_name, r, item_id, f"answerがa/b/c/d外: {answer!r}")
        else:
            selected = ws.cell(r, 4 + "abcd".index(answer)).value
            if cell_text(selected).strip() == "":
                append_issue(report, "needs_review", sheet_name, r, item_id, f"answer={answer} が指す選択肢が空欄")

        for c in range(4, 8):
            old = ws.cell(r, c).value
            if isinstance(old, str):
                new = old.strip()
                if new != old:
                    ws.cell(r, c).value = new
                    report["auto_fixes"].append({"sheet": sheet_name, "id": item_id, "kind": f"{EXPECTED_HEADERS[c-1]}前後空白除去", "before": old, "after": new})

        new_explanation = explanation
        applied = []
        for before, after in NG_REPLACEMENTS:
            if before in new_explanation:
                new_explanation = new_explanation.replace(before, after)
                applied.append(f"{before}->{after}")
        for before, after in REF_REPLACEMENTS:
            if before in new_explanation:
                new_explanation = new_explanation.replace(before, after)
                applied.append(f"{before}->{after}")
        if new_explanation != explanation:
            ws.cell(r, 9).value = new_explanation
            report["auto_fixes"].append({"sheet": sheet_name, "id": item_id, "kind": "解説NG/参照表現言い換え", "before": explanation, "after": new_explanation, "rules": applied})

        if question:
            if question in seen_questions:
                duplicate_rows.append(r)
                report["duplicates_within_block"].append({
                    "sheet": sheet_name,
                    "kept_row": seen_questions[question],
                    "deleted_row": r,
                    "question": question,
                })
            else:
                seen_questions[question] = r

        if any(term in question for term in AMBIGUOUS_TERMS):
            append_issue(report, "needs_review", sheet_name, r, item_id, f"曖昧表現を含む: {', '.join(t for t in AMBIGUOUS_TERMS if t in question)}")
        if 0 < len(question) < 15:
            append_issue(report, "needs_review", sheet_name, r, item_id, f"question15字未満: {question}")

        stripped_choices = [ch.strip() for ch in choices if ch.strip()]
        pair_counts = defaultdict(list)
        for idx, ch in enumerate([cell_text(ws.cell(r, c).value).strip() for c in range(4, 8)], start=0):
            if ch:
                pair_counts[ch].append("abcd"[idx])
        for ch, labels in pair_counts.items():
            if len(labels) > 1:
                append_issue(report, "needs_review", sheet_name, r, item_id, f"選択肢完全一致: {','.join(labels)} = {ch}")
        if stripped_choices:
            min_len = min(len(ch) for ch in stripped_choices if len(ch) > 0)
            max_len = max(len(ch) for ch in stripped_choices)
            if min_len > 0 and max_len / min_len >= 3:
                append_issue(report, "needs_review", sheet_name, r, item_id, f"選択肢文字数比 {max_len}/{min_len} = {max_len/min_len:.1f}倍")

    for r in sorted(duplicate_rows, reverse=True):
        ws.delete_rows(r, 1)
        report["auto_fixes"].append({"sheet": sheet_name, "id": None, "kind": "ブロック内question完全一致重複削除", "row": r})

    if duplicate_rows:
        compact_ids(ws)
        report["auto_fixes"].append({"sheet": sheet_name, "id": None, "kind": "id詰め直し", "count_deleted": len(duplicate_rows)})

    after = len(data_rows(ws))
    report["block_summary"][sheet_name]["after"] = after


def remove_cross_block_duplicates(wb, sheet_names, report):
    seen = {}
    to_delete = defaultdict(list)
    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        for r in range(2, ws.max_row + 1):
            if not any(ws.cell(r, c).value is not None for c in range(1, 10)):
                continue
            q = norm_question(ws.cell(r, 3).value)
            if not q:
                continue
            if q in seen:
                to_delete[sheet_name].append(r)
                report["duplicates_cross_block_deleted"].append({
                    "question": q,
                    "kept": {"sheet": seen[q][0], "row": seen[q][1], "id": seen[q][2]},
                    "deleted": {"sheet": sheet_name, "row": r, "id": ws.cell(r, 1).value},
                })
            else:
                seen[q] = (sheet_name, r, ws.cell(r, 1).value)

    for sheet_name, rows in to_delete.items():
        ws = wb[sheet_name]
        for r in sorted(rows, reverse=True):
            ws.delete_rows(r, 1)
            report["auto_fixes"].append({"sheet": sheet_name, "id": None, "kind": "ブロックまたぎquestion完全一致重複削除", "row": r})
        compact_ids(ws)
        report["auto_fixes"].append({"sheet": sheet_name, "id": None, "kind": "id詰め直し", "count_deleted": len(rows)})
        report["block_summary"][sheet_name]["after"] = len(data_rows(ws))


def collect_prefix_similarity(wb, sheet_names, report):
    buckets = defaultdict(list)
    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        for r in range(2, ws.max_row + 1):
            if not any(ws.cell(r, c).value is not None for c in range(1, 10)):
                continue
            prefix = question_prefix(ws.cell(r, 3).value)
            if len(prefix) == 30:
                buckets[prefix].append({"sheet": sheet_name, "row": r, "id": ws.cell(r, 1).value})
    for prefix, items in buckets.items():
        if len(items) > 1:
            report["prefix30_similar"].append({"prefix": prefix, "items": items})


def rebuild_final_format(wb, sheet_names, report):
    if "最終フォーマット" not in wb.sheetnames:
        ws = wb.create_sheet("最終フォーマット", 0)
        for c, h in enumerate(EXPECTED_HEADERS, start=1):
            ws.cell(1, c).value = h
    target = wb["最終フォーマット"]

    for r in range(2, target.max_row + 1):
        for c in range(1, 10):
            target.cell(r, c).value = None

    next_id = 1
    tr = 2
    seen = set()
    skipped = []
    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        for r in range(2, ws.max_row + 1):
            if not any(ws.cell(r, c).value is not None for c in range(1, 10)):
                continue
            q = norm_question(ws.cell(r, 3).value)
            if q in seen:
                skipped.append({"sheet": sheet_name, "row": r, "id": ws.cell(r, 1).value, "question": q})
                continue
            seen.add(q)
            values = [ws.cell(r, c).value for c in range(1, 10)]
            values[0] = next_id
            for c, value in enumerate(values, start=1):
                target.cell(tr, c).value = value
            next_id += 1
            tr += 1

    ids = [target.cell(r, 1).value for r in range(2, tr)]
    report["final_format"] = {
        "count": len(ids),
        "ids_sequential": ids == list(range(1, len(ids) + 1)),
        "skipped_cross_duplicates": skipped,
    }


def inspect_app_sheet(wb, report):
    if "アプリ内容" not in wb.sheetnames:
        report["app_content"]["exists"] = False
        return
    ws = wb["アプリ内容"]
    report["app_content"]["exists"] = True
    found = {}
    for r in range(2, ws.max_row + 1):
        item = cell_text(ws.cell(r, 1).value).strip()
        value = cell_text(ws.cell(r, 2).value)
        if not item:
            continue
        actual_len = len(value)
        current_len = ws.cell(r, 4).value
        found[item] = value
        if item in APP_LIMITS and actual_len > APP_LIMITS[item]:
            report["app_content"]["length_violations"].append({"item": item, "limit": APP_LIMITS[item], "actual": actual_len})
        if current_len != actual_len:
            report["app_content"]["current_length_mismatch"].append({"item": item, "stored": current_len, "actual": actual_len})
        for ng, _ in NG_REPLACEMENTS:
            if ng in value:
                report["app_content"]["ng_expressions"].append({"item": item, "expression": ng})
    for required in APP_REQUIRED:
        if not found.get(required, "").strip():
            report["app_content"]["missing"].append(required)
    report["app_content"]["filled_count"] = sum(1 for item in APP_REQUIRED if found.get(item, "").strip())


def main():
    wb = openpyxl.load_workbook(XLSX_PATH)
    sheet_names = [name for name in wb.sheetnames if name.startswith(SHEET_PREFIXES)]
    sheet_names.sort(key=lambda name: "①②③④⑤⑥⑦⑧⑨⑩".index(name[0]))

    report = {
        "target_file": str(XLSX_PATH),
        "verified_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "block_summary": {},
        "auto_fixes": [],
        "needs_review": [],
        "duplicates_within_block": [],
        "duplicates_cross_block_deleted": [],
        "prefix30_similar": [],
        "final_format": {},
        "app_content": {
            "exists": False,
            "filled_count": 0,
            "missing": [],
            "length_violations": [],
            "current_length_mismatch": [],
            "ng_expressions": [],
        },
    }

    for sheet_name in sheet_names:
        inspect_block_sheet(wb, sheet_name, report)

    remove_cross_block_duplicates(wb, sheet_names, report)
    collect_prefix_similarity(wb, sheet_names, report)
    rebuild_final_format(wb, sheet_names, report)
    inspect_app_sheet(wb, report)

    sheet_fix_counts = defaultdict(int)
    for fix in report["auto_fixes"]:
        if fix.get("sheet") in report["block_summary"]:
            sheet_fix_counts[fix["sheet"]] += 1
    sheet_review_counts = defaultdict(int)
    for issue in report["needs_review"]:
        if issue.get("sheet") in report["block_summary"]:
            sheet_review_counts[issue["sheet"]] += 1
    for sheet_name in sheet_names:
        report["block_summary"][sheet_name]["auto_fixes"] = sheet_fix_counts[sheet_name]
        report["block_summary"][sheet_name]["needs_review"] = sheet_review_counts[sheet_name]

    wb.save(XLSX_PATH)
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "saved": str(XLSX_PATH),
        "report": str(REPORT_PATH),
        "final_count": report["final_format"]["count"],
        "ids_sequential": report["final_format"]["ids_sequential"],
        "auto_fixes": len(report["auto_fixes"]),
        "needs_review": len(report["needs_review"]),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
