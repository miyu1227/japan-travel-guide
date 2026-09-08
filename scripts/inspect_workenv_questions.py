import json
import re
import sys

import openpyxl


path = r"C:\Users\miyul\Downloads\作業環境測定士.xlsx"
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
keys = [
    "管理濃度",
    "測定",
    "保存",
    "mSv",
    "Sv",
    "年",
    "月",
    "週",
    "μ",
    "ppm",
    "mg/m3",
    "cm3",
    "粒径",
    "半減期",
    "波長",
    "nm",
    "許容濃度",
    "第",
    "有機溶剤",
    "特定化学物質",
    "粉じん",
    "放射",
    "線量",
    "管理区分",
    "A測定",
    "B測定",
    "C測定",
    "D測定",
    "個人サンプリング",
    "吸光",
    "原子吸光",
    "ICP",
    "GC",
    "HPLC",
    "X線",
    "ガラス",
    "石英",
    "質量数",
    "不偏分散",
]

rows = []
for sheet_name in [name for name in wb.sheetnames if name[0] in "①②③④⑤⑥⑦⑧⑨"]:
    ws = wb[sheet_name]
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 10)]
        text = " ".join(str(v) for v in vals if v is not None)
        if any(k in text for k in keys) or re.search(r"\d", text):
            rows.append(
                {
                    "sheet": sheet_name,
                    "id": vals[0],
                    "answer": vals[7],
                    "question": vals[2],
                    "a": vals[3],
                    "b": vals[4],
                    "c": vals[5],
                    "d": vals[6],
                    "explanation": vals[8],
                }
            )

json.dump(rows, sys.stdout, ensure_ascii=False, indent=2)
