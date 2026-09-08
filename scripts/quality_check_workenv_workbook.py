from __future__ import annotations

import copy
import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl


INPUT = Path(r"C:\Users\miyul\Downloads\作業環境測定士.xlsx")
BACKUP = Path(r"C:\Users\miyul\Downloads\作業環境測定士.backup.xlsx")
QC_BACKUP = Path(r"C:\Users\miyul\Downloads\作業環境測定士.quality_check_backup.xlsx")
REPORT_JSON = Path(r"C:\Users\miyul\japan-travel-guide\workenv_quality_check_report.json")

BLOCK_SHEETS = [
    "① 労働衛生一般（共通）",
    "② 労働衛生関係法令（共通）",
    "③ デザイン・サンプリング（共通）",
    "④ 分析に関する概論（共通）",
    "⑤ 有機溶剤（第1種）",
    "⑥ 鉱物性粉じん（第1種）",
    "⑦ 特定化学物質（第1種）",
    "⑧ 金属類（第1種）",
    "⑨ 放射性物質（第1種）",
]
FINAL_SHEET = "最終フォーマット"
APP_SHEET = "アプリ内容"
HEADERS = ["id", "block", "question", "a", "b", "c", "d", "answer", "explanation"]
COL = {name: idx + 1 for idx, name in enumerate(HEADERS)}

AMBIGUOUS_TERMS = ["最も適切", "一般的に", "主に", "通常"]
TEXT_REF_TERMS = ["テキストでは", "本書に", "下表", "上図"]
NG_TERMS = ["絶対", "100%", "必ず", "確実", "保証"]
APP_FORBIDDEN = ["絶対合格", "100%合格", "合格保証"]
APP_LIMITS = {
    "アプリ名": 30,
    "サブタイトル": 30,
    "プロモーション用テキスト": 170,
    "概要": 4000,
    "キーワード": 100,
}


def value(ws, row: int, name: str):
    return ws.cell(row=row, column=COL[name]).value


def set_value(ws, row: int, name: str, val):
    ws.cell(row=row, column=COL[name]).value = val


def text(v) -> str:
    return "" if v is None else str(v)


def normalize_answer(v):
    if v is None:
        return None
    s = str(v).strip()
    trans = str.maketrans({"Ａ": "A", "Ｂ": "B", "Ｃ": "C", "Ｄ": "D", "ａ": "a", "ｂ": "b", "ｃ": "c", "ｄ": "d"})
    return s.translate(trans).lower()


def rewrite_question(q: str) -> str:
    out = q
    out = out.replace("テキストでは", "一般的な試験対策上、")
    out = out.replace("本書に", "問題文に")
    out = out.replace("下表", "次の条件")
    out = out.replace("上図", "問題文中の条件")
    return out


def rewrite_explanation(exp: str) -> str:
    out = exp
    out = out.replace("絶対", "原則として")
    out = out.replace("必ず", "原則として")
    out = out.replace("確実", "高い精度で")
    out = out.replace("保証", "担保")
    if "100%" in out and not re.search(r"(捕集率|捕集効率|計数効率|濃度限度|条件|計算|100\s*%)", out):
        out = out.replace("100%", "非常に高い割合")
    return out


def same_prefix_30(q: str) -> str:
    return q[:30]


def copy_cell_style(src, dst):
    if src.has_style:
        dst._style = copy.copy(src._style)
    if src.number_format:
        dst.number_format = src.number_format
    if src.alignment:
        dst.alignment = copy.copy(src.alignment)


def nonempty_block_rows(ws):
    return [r for r in range(2, ws.max_row + 1) if text(value(ws, r, "question")).strip()]


def trim_choice(s):
    return s.strip() if isinstance(s, str) else s


def check_app_sheet(ws):
    issues = {
        "filled_count": 0,
        "blank_items": [],
        "length_violations": [],
        "forbidden": [],
        "char_count_mismatch": [],
    }
    for row in range(2, ws.max_row + 1):
        item = text(ws.cell(row=row, column=1).value).strip()
        content = ws.cell(row=row, column=2).value
        actual_len = len(text(content))
        current = ws.cell(row=row, column=4).value
        if content is None or text(content).strip() == "":
            issues["blank_items"].append(item or f"row{row}")
        else:
            issues["filled_count"] += 1
        if item in APP_LIMITS and actual_len > APP_LIMITS[item]:
            issues["length_violations"].append({"item": item, "actual": actual_len, "limit": APP_LIMITS[item]})
        for phrase in APP_FORBIDDEN:
            if phrase in text(content):
                issues["forbidden"].append({"item": item, "phrase": phrase})
        if isinstance(current, int) and current != actual_len:
            issues["char_count_mismatch"].append({"item": item, "current": current, "actual": actual_len})
    return issues


def main():
    if not BACKUP.exists():
        shutil.copy2(INPUT, BACKUP)
    if not QC_BACKUP.exists():
        shutil.copy2(INPUT, QC_BACKUP)

    wb = openpyxl.load_workbook(INPUT)
    stats = {}
    auto_fixes = []
    user_checks = []
    cross_block_deleted = []

    for sheet_name in BLOCK_SHEETS:
        ws = wb[sheet_name]
        before = len(nonempty_block_rows(ws))
        stats[sheet_name] = {"before": before, "after": None, "auto": 0, "check": 0}

        seen_questions = {}
        rows_to_delete = []
        prefix_seen = defaultdict(list)

        for row in nonempty_block_rows(ws):
            q = text(value(ws, row, "question")).strip()
            if q in seen_questions:
                rows_to_delete.append(row)
                auto_fixes.append({"sheet": sheet_name, "row": row, "type": "block_duplicate_deleted", "detail": q[:80]})
                continue
            seen_questions[q] = row
            prefix_seen[same_prefix_30(q)].append(row)

            original_answer = value(ws, row, "answer")
            normalized = normalize_answer(original_answer)
            if normalized != original_answer:
                set_value(ws, row, "answer", normalized)
                auto_fixes.append({"sheet": sheet_name, "row": row, "type": "answer_normalized", "before": original_answer, "after": normalized})
            if normalized not in ["a", "b", "c", "d"]:
                user_checks.append({"sheet": sheet_name, "row": row, "type": "invalid_answer", "detail": original_answer})
            elif text(value(ws, row, normalized)).strip() == "":
                user_checks.append({"sheet": sheet_name, "row": row, "type": "answer_points_blank_choice", "detail": normalized})

            for choice in ["a", "b", "c", "d"]:
                old = value(ws, row, choice)
                new = trim_choice(old)
                if new != old:
                    set_value(ws, row, choice, new)
                    auto_fixes.append({"sheet": sheet_name, "row": row, "type": "choice_trim", "choice": choice})

            q2 = text(value(ws, row, "question"))
            if any(term in q2 for term in TEXT_REF_TERMS):
                new_q = rewrite_question(q2)
                set_value(ws, row, "question", new_q)
                auto_fixes.append({"sheet": sheet_name, "row": row, "type": "text_reference_rewritten", "before": q2, "after": new_q})

            exp = text(value(ws, row, "explanation"))
            matched_ng = [term for term in NG_TERMS if term in exp]
            if matched_ng:
                new_exp = rewrite_explanation(exp)
                if new_exp != exp:
                    set_value(ws, row, "explanation", new_exp)
                    auto_fixes.append({"sheet": sheet_name, "row": row, "type": "ng_expression_rewritten", "terms": matched_ng})
                else:
                    user_checks.append({"sheet": sheet_name, "row": row, "type": "ng_expression_left_as_condition", "terms": matched_ng})

            q_current = text(value(ws, row, "question"))
            amb = [term for term in AMBIGUOUS_TERMS if term in q_current]
            if amb:
                user_checks.append({"sheet": sheet_name, "row": row, "type": "ambiguous_expression", "terms": amb, "question": q_current[:80]})
            if len(q_current.strip()) < 15:
                user_checks.append({"sheet": sheet_name, "row": row, "type": "question_too_short", "question": q_current})

            choices = {c: text(value(ws, row, c)).strip() for c in ["a", "b", "c", "d"]}
            inv = defaultdict(list)
            for c, val in choices.items():
                if val:
                    inv[val].append(c)
            for val, labels in inv.items():
                if len(labels) >= 2:
                    user_checks.append({"sheet": sheet_name, "row": row, "type": "identical_choices", "choices": labels, "detail": val[:80]})
            lengths = [len(v) for v in choices.values() if v]
            if lengths and min(lengths) > 0 and max(lengths) / min(lengths) >= 3:
                user_checks.append({"sheet": sheet_name, "row": row, "type": "choice_length_ratio_ge_3", "min": min(lengths), "max": max(lengths), "question": q_current[:80]})

        for prefix, rows in prefix_seen.items():
            if prefix and len(rows) >= 2:
                user_checks.append({"sheet": sheet_name, "type": "same_first_30_chars", "prefix": prefix, "rows": rows})

        for row in sorted(rows_to_delete, reverse=True):
            ws.delete_rows(row, 1)

        # Re-number block ids after duplicate deletion.
        for idx, row in enumerate(nonempty_block_rows(ws), start=1):
            if value(ws, row, "id") != idx:
                set_value(ws, row, "id", idx)
                auto_fixes.append({"sheet": sheet_name, "row": row, "type": "block_id_renumbered", "after": idx})

        after = len(nonempty_block_rows(ws))
        stats[sheet_name]["after"] = after

    # Cross-block aggregation and exact duplicate deletion.
    final_ws = wb[FINAL_SHEET]
    for row in final_ws.iter_rows():
        for cell in row:
            cell.value = None
    for col_idx, header in enumerate(HEADERS, start=1):
        final_ws.cell(row=1, column=col_idx).value = header

    global_seen = {}
    out_row = 2
    global_id = 1
    for sheet_name in BLOCK_SHEETS:
        ws = wb[sheet_name]
        for row in nonempty_block_rows(ws):
            q = text(value(ws, row, "question")).strip()
            if q in global_seen:
                cross_block_deleted.append({"kept": global_seen[q], "deleted": {"sheet": sheet_name, "row": row}, "question": q[:100]})
                continue
            global_seen[q] = {"sheet": sheet_name, "row": row}
            for col_idx in range(1, len(HEADERS) + 1):
                dst = final_ws.cell(row=out_row, column=col_idx)
                src = ws.cell(row=row, column=col_idx)
                dst.value = global_id if col_idx == 1 else src.value
                copy_cell_style(src, dst)
            out_row += 1
            global_id += 1

    for col_idx in range(1, len(HEADERS) + 1):
        final_ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 12 if col_idx in [1, 2, 8] else 36

    final_ids = [final_ws.cell(row=r, column=1).value for r in range(2, out_row)]
    final_ok = final_ids == list(range(1, len(final_ids) + 1))
    app_report = check_app_sheet(wb[APP_SHEET])

    # Count auto/user by sheet.
    for fix in auto_fixes:
        if fix.get("sheet") in stats:
            stats[fix["sheet"]]["auto"] += 1
    for item in user_checks:
        if item.get("sheet") in stats:
            stats[item["sheet"]]["check"] += 1

    wb.save(INPUT)

    report = {
        "target": str(INPUT),
        "backup": str(BACKUP),
        "quality_backup": str(QC_BACKUP),
        "timestamp": datetime.now().isoformat(timespec="seconds"),
        "stats": stats,
        "auto_fixes": auto_fixes,
        "cross_block_deleted": cross_block_deleted,
        "user_checks": user_checks,
        "final_format": {"count": len(final_ids), "id_sequence_ok": final_ok},
        "app_sheet": app_report,
    }
    REPORT_JSON.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(REPORT_JSON)


if __name__ == "__main__":
    main()
